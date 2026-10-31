#!/usr/bin/env python3
"""Generate tests/fixtures/features.docx with python-docx (Word's default template).

Covers what the reader has to resolve in Word-made files: basedOn style chains,
lists driven by paragraph styles (List Bullet / List Number), a right tab stop
exactly on the margin, a left-border callout, a table with horizontal and
vertical merges and shading, an inline image, an external hyperlink, and a
page break.

Usage: python3 tools/make_fixture.py [output.docx]
Requires: pip install python-docx pillow
"""
import io
import sys

import docx
from docx.enum.text import WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def left_border(paragraph, color="2E75B6"):
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for k, v in (("w:val", "single"), ("w:sz", "12"), ("w:space", "8"), ("w:color", color)):
        left.set(qn(k), v)
    bdr.append(left)
    # pBdr must come before spacing/ind/jc in pPr
    ppr.insert(0, bdr)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def main(out):
    d = docx.Document()
    # 0.55in margins (792 twips = 52.8px): fractional-pixel margins are what
    # exposed the right-tab-on-margin wrap, so the fixture keeps them.
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Inches(0.55)
    d.add_heading("Feature Fixture", level=0)
    d.add_heading("Styles and runs", level=1)
    p = d.add_paragraph("Plain, ")
    p.add_run("bold").bold = True
    p.add_run(", ")
    r = p.add_run("italic colored")
    r.italic = True
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    p.add_run(", ")
    r = p.add_run("14pt Arial")
    r.font.size = Pt(14)
    r.font.name = "Arial"
    p.add_run(".")

    section = d.sections[0]
    content = section.page_width - section.left_margin - section.right_margin
    p = d.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(content, WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("Role at Company").bold = True
    p.add_run("\t2024 – Present").bold = True

    callout = d.add_paragraph("A left-border callout paragraph, indented and italic.")
    callout.paragraph_format.left_indent = Inches(0.25)
    callout.runs[0].italic = True
    left_border(callout)

    d.add_heading("Lists", level=1)
    for t in ("First bullet", "Second bullet"):
        d.add_paragraph(t, style="List Bullet")
    d.add_paragraph("Nested bullet", style="List Bullet 2")
    for t in ("Step one", "Step two", "Step three"):
        d.add_paragraph(t, style="List Number")

    d.add_heading("Table", level=1)
    t = d.add_table(rows=3, cols=3)
    t.style = "Table Grid"
    t.cell(0, 0).merge(t.cell(0, 1)).text = "Merged header"
    t.cell(0, 2).text = "H3"
    shade(t.cell(0, 2), "DDEBF7")
    t.cell(1, 0).text = "a1"
    t.cell(1, 1).text = "b1"
    t.cell(1, 2).merge(t.cell(2, 2)).text = "tall"
    t.cell(2, 0).text = "a2"
    t.cell(2, 1).text = "b2"

    d.add_heading("Image and link", level=1)
    img = Image.new("RGB", (240, 120), (0x33, 0x66, 0x99))
    ImageDraw.Draw(img).rectangle((20, 20, 100, 100), fill="white")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    d.add_picture(buf, width=Inches(2.5))
    p = d.add_paragraph("See ")
    add_hyperlink(p, "https://github.com/scottpeterman/anyprocessorqt", "the repo")
    p.add_run(" for details.")

    p = d.add_paragraph("Last line on page one.")
    p.add_run().add_break(WD_BREAK.PAGE)
    d.add_paragraph("First line on page two.")
    d.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "tests/fixtures/features.docx")
