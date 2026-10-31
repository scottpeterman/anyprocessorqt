#!/usr/bin/env python3
"""Compare two .docx files paragraph by paragraph with python-docx.

Checks text (including tabs and breaks), paragraph style, list membership,
image count, and per-run bold/italic/underline/size/font/color after
resolving direct formatting, character styles and paragraph styles.

Usage: python3 tools/compare_docx.py original.docx roundtrip.docx
Exit status is non-zero when differences are found.
"""
import sys

import docx
from docx.oxml.ns import qn


def style_chain_value(style, getter):
    while style is not None:
        v = getter(style)
        if v is not None:
            return v
        style = style.base_style
    return None


def doc_default_rpr(document):
    styles = document.styles.element
    rpr = styles.find(qn("w:docDefaults"))
    if rpr is None:
        return None
    rpr = rpr.find(qn("w:rPrDefault"))
    return None if rpr is None else rpr.find(qn("w:rPr"))


def default_attr(rpr, tag, attr="w:val"):
    if rpr is None:
        return None
    el = rpr.find(qn(tag))
    return None if el is None else el.get(qn(attr))


def effective_run(run, para, document):
    def resolve(name):
        v = getattr(run.font, name)
        if v is not None:
            return v
        if run.style is not None and run.style.type == 2:  # character style
            v = style_chain_value(run.style, lambda s: getattr(s.font, name))
            if v is not None:
                return v
        if para.style is None:
            return None
        return style_chain_value(para.style, lambda s: getattr(s.font, name))

    d = doc_default_rpr(document)
    bold = bool(resolve("bold"))
    italic = bool(resolve("italic"))
    underline = bool(resolve("underline"))
    size = resolve("size")
    if size is None:
        hp = default_attr(d, "w:sz")
        size = int(hp) * 6350 if hp else 127000
    name = resolve("name")
    if name is None:
        name = default_attr(d, "w:rFonts", "w:ascii") or "Times New Roman"
    color = None
    for src in (run.font, run.style.font if run.style is not None and run.style.type == 2 else None):
        if src is not None and src.color is not None and src.color.type is not None and src.color.rgb is not None:
            color = str(src.color.rgb)
            break
    if color is None and para.style is not None:
        color = style_chain_value(
            para.style,
            lambda s: str(s.font.color.rgb) if s.font.color is not None and s.font.color.type is not None
            and s.font.color.rgb is not None else None)
    return (bold, italic, underline, int(size), name, color or "auto")


def para_numbered(p):
    ppr = p._p.pPr
    if ppr is not None and ppr.numPr is not None:
        nid = ppr.numPr.numId
        return nid is not None and nid.val != 0
    if p.style is None:
        return False
    return style_chain_value(
        p.style, lambda s: True if s.element.pPr is not None and s.element.pPr.numPr is not None else None) or False


def runs_signature(p, document):
    # Merge adjacent runs with identical effective formatting; the writer may split/join runs.
    sig = []
    for r in p.runs:
        if not r.text:
            continue
        fmt = effective_run(r, p, document)
        if sig and sig[-1][1] == fmt:
            sig[-1] = (sig[-1][0] + r.text, fmt)
        else:
            sig.append((r.text, fmt))
    return sig


def all_paragraphs(document):
    out = []

    def walk(container):
        for block in container.iter_inner_content():
            if isinstance(block, docx.text.paragraph.Paragraph):
                out.append(block)
            else:
                for row in block.rows:
                    seen = set()
                    for cell in row.cells:
                        if id(cell._tc) in seen:
                            continue
                        seen.add(id(cell._tc))
                        walk(cell)

    walk(document)
    return out


def main(a_path, b_path):
    a, b = docx.Document(a_path), docx.Document(b_path)
    pa, pb = all_paragraphs(a), all_paragraphs(b)
    problems = 0
    if len(pa) != len(pb):
        print(f"paragraph count differs: {len(pa)} vs {len(pb)}")
        problems += 1
    for i, (x, y) in enumerate(zip(pa, pb)):
        if x.text != y.text:
            print(f"[{i}] text differs:\n  {x.text[:100]!r}\n  {y.text[:100]!r}")
            problems += 1
            continue
        sx_name = x.style.name if x.style is not None else None
        sy_name = y.style.name if y.style is not None else None
        if sx_name != sy_name:
            print(f"[{i}] style {sx_name!r} vs {sy_name!r}")
            problems += 1
        if para_numbered(x) != para_numbered(y):
            print(f"[{i}] list membership {para_numbered(x)} vs {para_numbered(y)}: {x.text[:60]!r}")
            problems += 1
        sx, sy = runs_signature(x, a), runs_signature(y, b)
        if sx != sy:
            print(f"[{i}] run formatting differs: {x.text[:60]!r}")
            for (tx, fx), (ty, fy) in zip(sx, sy):
                if (tx, fx) != (ty, fy):
                    print(f"     {tx[:30]!r} {fx}\n     {ty[:30]!r} {fy}")
                    break
            problems += 1
    ia = len(a.inline_shapes)
    ib = len(b.inline_shapes)
    if ia != ib:
        print(f"inline image count {ia} vs {ib}")
        problems += 1
    ta, tb = len(a.tables), len(b.tables)
    if ta != tb:
        print(f"table count {ta} vs {tb}")
        problems += 1
    print(f"{len(pa)} paragraphs, {ia} images, {ta} tables compared: {problems} difference(s)")
    return 1 if problems else 0


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(sys.argv[1], sys.argv[2]))
