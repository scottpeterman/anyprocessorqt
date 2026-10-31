#!/usr/bin/env python3
"""Generate tests/fixtures/headers.docx with python-docx (Word's default template).

Covers headers and footers: a default header with a complex PAGE / NUMPAGES
field pair on the Header style's right tab, a different first-page header,
a default footer with a simple PAGE field, and a body long enough for three
pages.

Usage: python3 tools/make_headers_fixture.py [output.docx]
Requires: pip install python-docx
"""
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def complex_field(paragraph, instr, result):
    def run_with(child):
        r = OxmlElement("w:r")
        r.append(child)
        paragraph._p.append(r)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run_with(begin)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    run_with(it)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run_with(sep)
    paragraph.add_run(result)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_with(end)


def simple_field(paragraph, instr, result):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), f" {instr} ")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = result
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


def main(out):
    d = docx.Document()
    sec = d.sections[0]
    sec.different_first_page_header_footer = True

    hp = sec.header.paragraphs[0]
    hp.add_run("Fixture Report\t\tPage ")
    complex_field(hp, "PAGE", "1")
    hp.add_run(" of ")
    complex_field(hp, "NUMPAGES", "3")

    sec.first_page_header.paragraphs[0].add_run("FIRST PAGE HEADER")

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("- ")
    simple_field(fp, "PAGE", "1")
    fp.add_run(" -")

    d.add_heading("Headers and footers", 1)
    for i in range(70):
        d.add_paragraph(
            f"Body paragraph {i}: enough text to wrap across the line so that three pages "
            "fill up with ordinary paragraphs and the header and footer repeat on each."
        )
    d.save(out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "tests/fixtures/headers.docx")
